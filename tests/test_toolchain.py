from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from scripts.quick_validate import CONFIG, make_fixture

from native_word_thesis.ooxml import NS, make_omath, node_text, polish_docx
from native_word_thesis.validate import validate_docx


def qn(prefix: str, tag: str) -> str:
    return f"{{{NS[prefix]}}}{tag}"


def test_polish_removes_common_conversion_artifacts() -> None:
    with tempfile.TemporaryDirectory(prefix="nwt-test-") as td:
        root = Path(td)
        fixture = root / "fixture.docx"
        output = root / "output.docx"
        make_fixture(fixture)
        polish_docx(fixture, output, CONFIG)
        report = validate_docx(output, CONFIG)
        with zipfile.ZipFile(output) as zf:
            document_root = etree.fromstring(zf.read("word/document.xml"))

    assert report["ok"]
    assert report["reference_numPr_left"] == 0
    assert report["math_object_count"] >= 1
    assert report["heading_italic_left"] == []
    assert report["uncentered_exact_texts"] == []
    assert report["cover_table_count"] == 1
    refs = [p for p in document_root.xpath("//w:body/w:p", namespaces=NS) if node_text(p).startswith("[1]")]
    assert refs and not refs[0].xpath("./w:pPr/w:wordWrap", namespaces=NS)
    assert refs[0].xpath("./w:pPr/w:suppressAutoHyphens[@w:val='1']", namespaces=NS)
    ref_text_runs = [
        "".join(run.xpath("./w:t/text()", namespaces=NS))
        for run in refs[0].xpath(".//w:r[not(ancestor::m:oMath)]", namespaces=NS)
    ]
    assert len([item for item in ref_text_runs if item]) == 1
    dates = [p for p in document_root.xpath("//w:body/w:p", namespaces=NS) if node_text(p) == "2026年6月"]
    assert dates and dates[0].xpath("./w:pPr/w:jc[@w:val='center']", namespaces=NS)


def test_omml_builder_supports_display_formula_shapes() -> None:
    formula = make_omath(
        [
            {"text": "CTR"},
            " = ",
            {"frac": [["1"], ["10"]]},
            {"sum": {"sub": ["o∈", {"text": "Top10"}], "body": [{"rad": [{"text": "rel"}, "(v,o)"]}]}},
        ]
    )

    assert node_text(formula) == "CTR=110o∈Top10rel(v,o)"
    assert len(formula.xpath(".//m:f", namespaces=NS)) == 1
    assert len(formula.xpath(".//m:nary", namespaces=NS)) == 1
    assert len(formula.xpath(".//m:rad", namespaces=NS)) == 1
    assert formula.xpath('.//m:r[m:rPr/m:nor]/w:rPr/w:rFonts[@w:ascii="Times New Roman"]', namespaces=NS)
    assert formula.xpath('.//m:r[not(m:rPr/m:nor)]/w:rPr/w:rFonts[@w:ascii="STIX Two Math"]', namespaces=NS)
    assert formula.xpath('.//m:r[m:t="="]/w:rPr/w:rFonts[@w:ascii="Times New Roman"]', namespaces=NS)
    assert not formula.xpath('.//m:r[m:t=" = "]', namespaces=NS)


def test_ooxml_guard_catches_field_font_and_numeric_artifacts() -> None:
    with tempfile.TemporaryDirectory(prefix="nwt-guard-test-") as td:
        root = Path(td)
        docx = root / "artifact.docx"
        cfg = root / "guard.json"

        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE

        doc = Document()
        compact_style = doc.styles.add_style("Compact", WD_STYLE_TYPE.PARAGRAPH)
        compact_style.name = "Compact"
        leaked = doc.add_paragraph("converted table/body residue")
        leaked.style = compact_style
        doc.add_paragraph('图 3.2 系统功能模块结构图 TC "图 3.2 系统功能模块结构图" \\f F \\l 1')
        bare_citation = doc.add_paragraph("引用缺少方括号")
        bare_run = bare_citation.add_run("12")
        bare_run.font.superscript = True
        latin = doc.add_paragraph().add_run("ASP.NET Core")
        latin.font.name = "Arial"
        math_para = doc.add_paragraph()
        math_para._p.append(make_omath(["100"]))
        missing_font_math = etree.Element(qn("m", "oMath"))
        missing_run = etree.SubElement(missing_font_math, qn("m", "r"))
        etree.SubElement(missing_run, qn("m", "t")).text = "x"
        math_para._p.append(missing_font_math)
        wrong_upright_math = etree.Element(qn("m", "oMath"))
        wrong_upright_run = etree.SubElement(wrong_upright_math, qn("m", "r"))
        wrong_upright_mrpr = etree.SubElement(wrong_upright_run, qn("m", "rPr"))
        etree.SubElement(wrong_upright_mrpr, qn("m", "nor"))
        wrong_upright_wrpr = etree.SubElement(wrong_upright_run, qn("w", "rPr"))
        wrong_upright_fonts = etree.SubElement(wrong_upright_wrpr, qn("w", "rFonts"))
        for attr in ["ascii", "hAnsi", "eastAsia", "cs"]:
            wrong_upright_fonts.set(qn("w", attr), "STIX Two Math")
        etree.SubElement(wrong_upright_run, qn("m", "t")).text = "Precision"
        math_para._p.append(wrong_upright_math)
        doc.save(docx)

        cfg.write_text(
            json.dumps(
                {
                    "heading_styles": [],
                    "forbid_heading4": False,
                    "caption_pattern": r"__never_match__",
                    "require_caption_center": False,
                    "require_caption_zero_indent": False,
                    "require_drawing_center": False,
                    "require_drawing_zero_indent": False,
                    "citations": {
                        "require_superscript_brackets": True,
                        "scan_before_reference_heading": False,
                    },
                    "field_codes": {
                        "forbid_instr_text_patterns": [r"\bTC\b", r"TOC\s+\\h\s+\\z\s+\\f"],
                        "forbid_visible_text_patterns": [r"\bTC\s+\"", r"\\f\s+[FT]\b", r"\\l\s+1\b"],
                    },
                    "latin_fonts": {
                        "required_ascii_hansi": "Times New Roman",
                        "require_direct_on_latin_runs": True,
                    },
                    "math": {
                        "forbid_simple_numeric_omml": True,
                        "required_math_font": "STIX Two Math",
                        "allowed_math_run_fonts": ["STIX Two Math", "Times New Roman"],
                        "upright_text_math_font": "Times New Roman",
                        "operator_math_font": "Times New Roman",
                        "forbid_spaced_operator_runs": True,
                        "forbidden_math_fonts": ["Cambria Math"],
                        "require_direct_math_run_font": True,
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        result = subprocess.run(
            [
                sys.executable,
                str(Path(__file__).parents[1] / "scripts/ooxml_thesis_guard.py"),
                str(docx),
                "--config",
                str(cfg),
                "--fail-on-warnings",
            ],
            check=False,
            capture_output=True,
            text=True,
        )

    assert result.returncode == 1
    report = json.loads(result.stdout)
    codes = {warning["code"] for warning in report["warnings"]}
    assert "forbidden_field_code_text" in codes
    assert "citation_superscript_missing_brackets" in codes
    assert "latin_font_mismatch" in codes
    assert "conversion_style_residue" in codes
    assert "math_run_font_missing" in codes
    assert "math_upright_text_font_mismatch" in codes
    assert "simple_numeric_omml" in codes
