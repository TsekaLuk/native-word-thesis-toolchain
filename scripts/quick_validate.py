#!/usr/bin/env python3
from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document

from native_word_thesis.ooxml import polish_docx
from native_word_thesis.validate import validate_docx


CONFIG = {
    "references": {"heading": "参考文献", "stop_prefixes": ["附录"]},
    "math_tables": [
        {
            "caption": "表A.1 推荐评分各分量权重与计算约定",
            "widths": [2100, 1450, 1300, 3900],
            "center_columns": [0, 1, 2],
            "nowrap_columns": [1, 2],
            "rows": [
                ["分量", "符号", "权重", "计算方式"],
                ["向量语义相似度", {"math": [{"sub": ["s", "vec"]}]}, "0.35", "余弦相似度"],
            ],
        }
    ],
}


def make_fixture(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("参考文献")
    ref = doc.add_paragraph("张三. 示例文献[J]. 示例学报, 2024.")
    ref.style = "List Number"
    doc.add_paragraph("附录 A 推荐算法权重表与接口规范")
    h3 = doc.add_paragraph()
    h3.style = "Heading 3"
    run = h3.add_run("A.1 混合推荐引擎评分权重")
    run.italic = True
    doc.add_paragraph("表A.1 推荐评分各分量权重与计算约定")
    table = doc.add_table(rows=2, cols=4)
    table.cell(0, 0).text = "分量"
    table.cell(0, 1).text = "符号"
    table.cell(0, 2).text = "权重"
    table.cell(0, 3).text = "计算方式"
    table.cell(1, 0).text = "向量语义相似度"
    table.cell(1, 1).text = "s_vec"
    table.cell(1, 2).text = "0.35"
    table.cell(1, 3).text = "余弦相似度"
    doc.save(path)


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="nwt-quick-") as td:
        root = Path(td)
        fixture = root / "fixture.docx"
        output = root / "output.docx"
        make_fixture(fixture)
        polish_docx(fixture, output, CONFIG)
        report = validate_docx(output, CONFIG)
        assert report["zip_ok"], report
        assert report["reference_count"] == 1, report
        assert report["reference_numPr_left"] == 0, report
        assert report["math_object_count"] >= 1, report
        assert not report["heading_italic_left"], report
        assert report["ok"], report
    print("quick_validate OK")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
